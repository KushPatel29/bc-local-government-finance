"""The workbook, held to the gold tables and to the pandas reference model.

Excel computes everything with formulas over the gold tables for the municipality
chosen on the Dashboard (Fernie in the committed file). These tests read the
saved workbook's cached values and recompute the same figures independently.
"""
from __future__ import annotations

import re
from datetime import datetime

import numpy as np
import pytest

import layout as L
import reference as R
from workbook_io import ROOT, block, cell, parts, power_query_m, spill

CENT = 0.01
TIGHT = 1e-9
ERRORS = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!", "#GETTING_DATA",
          "#FIELD!", "#BUSY!")
WORKBOOK = ROOT / "workbook" / L.WORKBOOK
CALC_SHEETS = ["Dashboard", "Profile", "Peers", "Assets", "Plan", "Tax", "Debt", "Province", "Checks"]


def test_every_check_passes(values):
    status = [r[0] for r in block(values, "checks_status")]
    assert status and all(s == "PASS" for s in status), status
    assert cell(values, "checks_summary") == f"All {len(status)} checks pass"


@pytest.mark.parametrize("sheet", CALC_SHEETS + ["Assumptions", "Cover"])
def test_no_cell_holds_an_error(values, sheet):
    bad = [(c.coordinate, c.value) for row in values[sheet].iter_rows() for c in row
           if isinstance(c.value, str) and c.value.startswith(ERRORS)]
    assert not bad, bad[:5]


def test_the_selection(values, gold, selected):
    assert cell(values, "Selected") == L.DEFAULT_MUNICIPALITY
    assert cell(values, "SelId") == selected
    assert cell(values, "BaseYear") == R.base_year(gold, selected)


def test_profile_is_the_gold_table(values, gold, selected):
    keys = [r[0] for r in block(values, "profile_keys")]
    years = block(values, "profile_years")[0]
    for key, row in zip(keys, block(values, "profile_block")):
        if key is None:
            continue
        for y, got in zip(years, row):
            want = R.value(gold, selected, key, int(y))
            if want is None:
                assert got in (None, ""), (key, y)
            else:
                assert got == pytest.approx(want, rel=1e-12, abs=1e-9), (key, y)


def test_peer_benchmark(values, gold, selected):
    peers = gold.peers.set_index("municipality_id").loc[selected]
    for row, (key, *_rest) in zip(block(values, "peer_benchmark"), L.METRICS):
        label, k, v, median, tier_pct, prov_pct, direction, reading = row
        assert k == key
        if np.isnan(peers[key]):
            assert v in (None, "") and reading == "not reported"
            continue
        assert v == pytest.approx(peers[key], rel=1e-12)
        assert median == pytest.approx(peers[f"{key}_tier_median"], rel=1e-12)
        assert tier_pct == pytest.approx(R.tier_percentile(gold, selected, key), abs=TIGHT)
        assert prov_pct == pytest.approx(peers[f"{key}_province_pct"], abs=TIGHT)


def test_the_tier_ranking_is_sorted_and_complete(values, gold, selected):
    rows = spill(values, "peer_ranking", 3)
    tier = gold.peers[gold.peers.size_tier == gold.peers.set_index("municipality_id").loc[selected, "size_tier"]]
    reporting = tier[tier.limit_used.notna()]
    assert len(rows) == len(reporting)
    vals = [r[2] for r in rows]
    assert vals == sorted(vals, reverse=True)
    assert L.DEFAULT_MUNICIPALITY in [r[0] for r in rows]


PLAN_KEYS = ["year", "cash_expenses", "non_tax_revenue", "capital", "incremental_capital", "new_borrowing",
             "new_servicing", "reserve_draw", "tax_capital", "taxation", "tax_increase", "reserves", "servicing",
             "limit_used", "house_tax"]


def test_plan_matches_the_reference_model(values, gold, selected):
    want = R.plan(gold, selected)
    for key, row in zip(PLAN_KEYS, block(values, "plan_table")):
        for t, got in enumerate(row):
            w = want[key].iloc[t]
            if key == "tax_increase" and t == 0:
                continue
            assert got == pytest.approx(w, rel=1e-10, abs=1e-6), (key, t)
    assert cell(values, "plan_avg_increase") == pytest.approx(R.average_increase(want), rel=1e-12)


def test_plan_data_table_re_runs_the_plan(values, gold, selected):
    want = R.plan_grid(gold, selected)
    for got_row, (_c, want_row) in zip(block(values, "plan_grid"), want.iterrows()):
        assert got_row == pytest.approx(list(want_row), rel=1e-10)


def test_levy_to_rates(values, gold, selected):
    change = cell(values, "tax_levy_change")
    want = R.levy_to_rate(gold, selected, change)
    assert change == pytest.approx(R.plan(gold, selected).tax_increase.iloc[1], rel=1e-12)
    assert cell(values, "tax_variable_levy") == pytest.approx(want["levy"], abs=CENT)
    assert cell(values, "tax_res_rate_implied") == pytest.approx(want["implied"], rel=1e-12)
    assert cell(values, "tax_res_rate_published") == pytest.approx(want["published"], rel=1e-12)
    assert want["implied"] == pytest.approx(want["published"], rel=1e-5), "the multiples reproduce the rate"
    assert cell(values, "tax_house_new") == pytest.approx(want["house_new"], rel=1e-12)


def test_borrowing_room(values, gold, selected):
    want = R.debt(gold, selected, amount=cell(values, "debt_amount"))
    assert cell(values, "debt_limit_used") == pytest.approx(want["limit_used"], rel=1e-12)
    assert cell(values, "debt_payment") == pytest.approx(want["payment"], rel=1e-10)
    assert cell(values, "debt_max_borrowing") == pytest.approx(want["max_borrowing"], rel=1e-10)


def test_assets_table_is_the_selection_sorted(values, gold, selected):
    rows = spill(values, "assets_table", 7)
    latest = gold.assets[(gold.assets.municipality_id == selected) & (gold.assets.year == gold.year.year.max())]
    assert len(rows) == len(latest)
    for r in rows:
        want = latest.set_index("asset_category").loc[r[0]]
        assert r[5] == pytest.approx(0 if np.isnan(want.replacement_value) else want.replacement_value, abs=CENT)
    assert [r[5] or 0 for r in rows] == sorted([r[5] or 0 for r in rows], reverse=True)


def test_province_facts(values, gold):
    t = gold.province.set_index("year").loc[gold.year.year.max()]
    assert cell(values, "prov_screened") == t.screened_for_renewal
    assert cell(values, "prov_below") == t.below_renewal_need
    assert cell(values, "prov_reinvestment") == pytest.approx(t.reinvestment_ratio, rel=1e-12)
    assert cell(values, "prov_missing") == 9


def test_power_query_outputs(values, gold):
    latest = gold.assets[gold.assets.year == gold.assets.year.max()]
    want = latest.groupby("asset_category").replacement_value.sum()
    ws = values["PQ_Categories"]
    head = [c.value for c in ws[4]]
    got = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[1] is None:
            break
        got[row[head.index("asset_category")]] = row[head.index("replacement_value")]
    assert set(got) == set(want.index)
    for k, v in got.items():
        assert v == pytest.approx(want[k], abs=1)
    tiers = gold.peers.groupby("size_tier").house_total_charges.median()
    ws = values["PQ_Tiers"]
    head = [c.value for c in ws[4]]
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[1] is None:
            break
        assert row[head.index("median_house_total_charges")] == pytest.approx(tiers[row[head.index("size_tier")]])


def test_dashboard_writes_its_own_sentences(values, gold, selected):
    first, second, third, fourth, fifth = [r[0] for r in block(values, "dash_commentary")]
    by = R.base_year(gold, selected)
    revenue = R.value(gold, selected, "total_revenue", by)
    assert f"${revenue / 1e6:,.1f}M" in first
    assert f"{R.average_increase(R.plan(gold, selected)):.1%}" in fourth
    assert f"${R.value(gold, selected, 'house_total_charges', gold.year.year.max()):,.0f}" in fifth
    assert "renewal need" in third and "limit" in second


# ------------------------------------------------------------ Excel features
@pytest.fixture(scope="module")
def pkg():
    return parts(WORKBOOK)


def test_power_query_is_real_m(pkg):
    m = power_query_m(pkg)
    assert "shared Categories" in m and "shared Tiers" in m
    assert "Table.Group" in m and "List.Median" in m


def test_named_lambdas(formulas):
    for name in ("YV", "MONEY"):
        assert formulas.defined_names[name].attr_text.startswith("_xlfn.LAMBDA("), name


def test_the_plan_grid_is_a_what_if_data_table(pkg, formulas):
    idx = formulas.sheetnames.index("Plan") + 1
    xml = pkg[f"xl/worksheets/sheet{idx}.xml"].decode()
    assert len(re.findall(r'<f t="dataTable"[^>]*dt2D="1"', xml)) == 1


def test_the_selector_is_a_validated_list(formulas):
    dv = [d for d in formulas["Dashboard"].data_validations.dataValidation if "C3" in str(d.sqref)]
    assert dv and dv[0].type == "list" and dv[0].formula1 == "ActiveList"


@pytest.mark.parametrize("function,sheet", [
    ("FILTER", "Assets"), ("SORTBY", "Peers"), ("XLOOKUP", "Peers"), ("LET", "Peers"), ("PMT", "Plan"),
    ("PV", "Debt"), ("CHOOSECOLS", "Tax"), ("HSTACK", "Tax"), ("MAP", "Checks"), ("MAXIFS", "Assumptions"),
])
def test_modern_functions_in_use(formulas, function, sheet):
    text = " ".join(str(getattr(c.value, "text", c.value)) for row in formulas[sheet].iter_rows() for c in row
                    if c.value is not None)
    assert function in text, f"{function} not on {sheet}"


def spilled(ws) -> set[str]:
    """Cells a dynamic array wrote: openpyxl reads them back as plain values."""
    from openpyxl.utils import get_column_letter, range_boundaries
    from openpyxl.worksheet.formula import ArrayFormula
    out = set()
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, ArrayFormula) and c.value.ref:
                c1, r1, c2, r2 = range_boundaries(c.value.ref)
                out |= {f"{get_column_letter(x)}{y}" for y in range(r1, r2 + 1) for x in range(c1, c2 + 1)}
    return out


def typed_numbers(ws, allowed):
    owned = spilled(ws) | set(allowed)
    return [c.coordinate for row in ws.iter_rows() for c in row
            if isinstance(c.value, (int, float, datetime)) and not isinstance(c.value, bool)
            and c.coordinate not in owned]


@pytest.mark.parametrize("sheet", ["Dashboard", "Profile", "Peers", "Assets", "Tax", "Province"])
def test_no_typed_numbers(formulas, sheet):
    allowed = {f"{c}5" for c in "DEFGHI"} if sheet == "Profile" else set()
    if sheet == "Province":
        allowed = {f"B{r}" for r in range(6, 12)}
    assert not typed_numbers(formulas[sheet], allowed)


def test_workbook_opens_on_the_cover_with_the_disclaimer(values):
    assert values.active.title == "Cover"
    text = " ".join(str(c.value) for row in values["Cover"].iter_rows() for c in row if c.value)
    assert "Not affiliated with or endorsed by any municipality or by the Province" in text
    assert "Open Government Licence" in text
